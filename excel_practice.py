import openpyxl as xl
wb = xl.load_workbook('computer.xlsx')
print('Type of wb:', type(wb))
print('Active worksheet object :', wb.active)
print('Worksheet names :', wb.sheetnames)
print('Worksheet objetcs :', wb.worksheets)
ws = wb['Sheet1']
print('Type of ws:', type(ws))

'''import openpyxl as xl
wb = xl.load_workbook(r'computer.xlsx')
ws = wb['Sheet1']
#returns the value of the cell 'A1'
cell = ws.cell(1, 1)
print(cell.value)
#returns the value of the cell 'A2'
cell = ws['A2']
print(cell.value)

import openpyxl as xl
wb = xl.load_workbook(r'computer.xlsx')
ws = wb['Sheet1']
#Adding as a list
ws.append( [1004, 'abc','def', 2020] )
#Adding as a tuple
ws.append( (1003, 'aaa','bbb', 2022) )
#Saving the changes
wb.save(r'computer.xlsx')

import openpyxl as xl
wb = xl.load_workbook(r'computer.xlsx')
ws = wb['Sheet1']
#Accessing the cell directly
ws['B4'] = 'James'
wb.save(r'computer.xlsx')'''




